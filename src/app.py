from flask import Flask, render_template, request, redirect, url_for, Response, jsonify
import database
import spider
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime  # 用于生成时间戳

app = Flask(__name__)

@app.route('/')
def index():
    keyword = request.args.get('keyword', '')
    msg = request.args.get('msg', '')
    jobs = database.get_all_jobs(keyword)
    return render_template('index.html', jobs=jobs, keyword=keyword, msg=msg)

@app.route('/crawl')
def crawl_data():
    count = spider.crawl_real_jobs()
    if count > 0:
        return redirect(url_for('index', msg=f"🎉 成功抓取并清洗了 {count} 条新数据（已包含中文翻译）！"))
    else:
        return redirect(url_for('index', msg="⚠️ 本次未抓到新数据，或网络受限。"))

@app.route('/delete/<int:job_id>')
def delete(job_id):
    """删除数据（AJAX 无刷新版）"""
    database.delete_job(job_id)
    # 返回 JSON 告诉前端删除成功，前端会直接移除那一行并重排序号
    return jsonify({"status": "success", "message": "删除成功"})

@app.route('/export')
def export_excel():
    jobs = database.get_all_jobs()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "招聘数据"

    # 表头
    headers = ['序号', '英文原版职位名', '中文翻译名', '公司名称', '薪资范围', '地点', '发布日期', 'URL']
    ws.append(headers)

    # 表头样式（蓝色背景，微软雅黑，白色加粗，居中）
    header_fill = PatternFill(start_color="4361ee", end_color="4361ee", fill_type="solid")
    header_font = Font(name='微软雅黑', color="FFFFFF", bold=True, size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 👇 核心修改点：遍历数据时自动生成 1, 2, 3... 的连续序号，忽略数据库原本断层的 ID
    for index, job in enumerate(jobs, start=1):
        # index 是连续序号，job[1:] 是除了 ID 以外的其他数据
        row_data = [index] + list(job[1:])
        ws.append(row_data)

    # 设置列宽（解决文字挤在一起的问题）
    column_widths = [8, 30, 30, 25, 15, 15, 12, 50]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 数据单元格样式（微软雅黑，垂直居中，自动换行）
    data_font = Font(name='微软雅黑', size=11)
    for row in ws.iter_rows(min_row=2, max_col=8):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = data_font

    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # 生成带时间戳的文件名，防止浏览器缓存旧文件
    current_time = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"jobs_data_{current_time}.xlsx"

    # 返回下载响应
    return Response(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}",
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"
        }
    )

if __name__ == '__main__':
    database.init_db()
    app.run(debug=True)